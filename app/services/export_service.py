import csv, io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask import Response, send_file
from app.models.schedule import ClassTimeSchedule

HDR_FILL = PatternFill(fill_type='solid', fgColor='1F3864')
HDR_FONT = Font(bold=True, color='FFFFFF')

HEADERS = ['Teacher','Course','Course Type','Room','Building',
           'Day','Start','End','Section','Program','Semester Label','Status']

def _rows():
    for s in ClassTimeSchedule.query.all():
        yield [
            s.teacher.full_name, s.course.title, s.course.course_type,
            s.room.room_type + s.room.room_number, s.room.building,
            s.time_slot.day, s.time_slot.start_time, s.time_slot.end_time,
            s.section.name, s.section.program, s.semester_label, s.status
        ]

def export_csv():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    for row in _rows(): w.writerow(row)
    buf.seek(0)
    return Response(buf.getvalue(), mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=timetable.csv'})

def export_excel():
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Timetable'
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal='center')
    for row in _rows(): ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name='timetable.xlsx')
